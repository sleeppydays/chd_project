import os
import json
import cv2
import numpy as np
import pandas as pd
import torch
import torch.nn as nn
from PIL import Image
from torch.utils.data import DataLoader, Dataset
from torchvision import transforms
from sklearn.model_selection import StratifiedGroupKFold
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from tqdm import tqdm

from pipeline_utils import (
    IMG_SIZE, IMAGENET_MEAN, IMAGENET_STD, get_model, get_target_layer,
    get_group_id, compute_binary_metrics, GradCAM, describe_focus_region, overlay_heatmap,
)

# ==========================================
# CONFIG — edit these
# ==========================================
# ==========================================
# CONFIG — STEP 3: ABLATION & GRAD-CAM
# ==========================================
REPORT_DIR = "/home/aaush/cardium_native"
MODEL_DIR = os.path.join(REPORT_DIR, "models")
HEATMAP_DIR = os.path.join(REPORT_DIR, "heatmaps")
MASTER_REPORT_PATH = os.path.join(REPORT_DIR, "ieee_master_report.xlsx")
CONDITIONS = {
    "A_raw": os.path.join(REPORT_DIR, "processed_A_raw"),
    "B_static": os.path.join(REPORT_DIR, "processed_B_static"),
    "C_adaptive": os.path.join(REPORT_DIR, "processed_C_adaptive"),
}
SEED = 42
EPOCHS = 8
BATCH_SIZE = 32
NAVY = "1F3864"
torch.backends.cudnn.benchmark = True
os.makedirs(HEATMAP_DIR, exist_ok=True)
device = "cuda" if torch.cuda.is_available() else "cpu"

with open(os.path.join(MODEL_DIR, "class_names.json")) as f:
    class_names = json.load(f)
with open(os.path.join(MODEL_DIR, "best_arch.json")) as f:
    best_info = json.load(f)
best_arch, POSITIVE_CLASS = best_info["architecture"], best_info["positive_class"]
class_to_idx = {c: i for i, c in enumerate(class_names)}
pos_idx = class_names.index(POSITIVE_CLASS)
print(f"Using architecture: {best_arch} (selected by script 2)")

tf = transforms.Compose([
    transforms.Resize((IMG_SIZE, IMG_SIZE)),
    transforms.ToTensor(),
    transforms.Normalize(mean=IMAGENET_MEAN, std=IMAGENET_STD),
])
train_tf = transforms.Compose([
    transforms.Resize((IMG_SIZE, IMG_SIZE)),
    transforms.RandomRotation(5),
    transforms.ColorJitter(brightness=0.15, contrast=0.15),
    transforms.ToTensor(),
    transforms.Normalize(mean=IMAGENET_MEAN, std=IMAGENET_STD),
    transforms.RandomErasing(p=0.5, scale=(0.02, 0.15)),
])


class CardiumDataset(Dataset):
    def __init__(self, samples, transform):
        self.samples = samples
        self.transform = transform

    def __len__(self):
        return len(self.samples)

    def __getitem__(self, idx):
        path, label = self.samples[idx]
        img = Image.open(path).convert("RGB")
        return self.transform(img), label


# ==========================================
# ONE ALIGNED SPLIT, REUSED ACROSS ALL 3 CONDITIONS
# Condition A/B/C directories mirror the same class/filename structure
# (built together in script 1), so we compute the split once from
# filenames and re-key it onto each condition's directory.
# ==========================================
master_list = []  # (class_name, filename)
for c in class_names:
    class_dir = os.path.join(CONDITIONS["A_raw"], c)
    for fn in sorted(os.listdir(class_dir)):
        if fn.lower().endswith((".png", ".jpg", ".jpeg")):
            master_list.append((c, fn))

labels = np.array([class_to_idx[c] for c, fn in master_list])
groups = np.array([get_group_id(fn) for c, fn in master_list])
gkf = StratifiedGroupKFold(n_splits=5, shuffle=True, random_state=SEED)
train_idx, val_idx = next(iter(gkf.split(np.zeros(len(master_list)), labels, groups)))
print(f"Ablation split: {len(train_idx)} train / {len(val_idx)} val (held fixed across A/B/C)")


def samples_for_condition(cond_key, indices):
    root = CONDITIONS[cond_key]
    out = []
    for i in indices:
        c, fn = master_list[i]
        out.append((os.path.join(root, c, fn), class_to_idx[c]))
    return out


def train_and_eval(samples_train, samples_val, arch):
    train_loader = DataLoader(CardiumDataset(samples_train, train_tf),
                               batch_size=BATCH_SIZE, shuffle=True, num_workers=4, pin_memory=True)
    val_loader = DataLoader(CardiumDataset(samples_val, tf),
                             batch_size=BATCH_SIZE, shuffle=False, num_workers=4, pin_memory=True)

    model = get_model(arch, len(class_names)).to(device)
    train_labels = [lbl for _, lbl in samples_train]
    counts = np.bincount(train_labels, minlength=len(class_names))
    weights = torch.tensor(1.0 / np.maximum(counts, 1), dtype=torch.float32)
    weights = (weights / weights.sum() * len(class_names)).to(device)
    criterion = nn.CrossEntropyLoss(weight=weights)
    optimizer = torch.optim.Adam(model.parameters(), lr=1e-4)

    for epoch in range(EPOCHS):
        model.train()
        for imgs, lbls in tqdm(train_loader, desc=f"Epoch {epoch+1}/{EPOCHS}"):
            imgs, lbls = imgs.to(device, non_blocking=True), lbls.to(device, non_blocking=True)
            optimizer.zero_grad()
            loss = criterion(model(imgs), lbls)
            loss.backward()
            optimizer.step()

    model.eval()
    all_labels, all_preds, all_scores = [], [], []
    with torch.no_grad():
        for imgs, lbls in val_loader:
            outputs = model(imgs.to(device))
            probs = torch.softmax(outputs, dim=1)
            all_preds.extend(outputs.argmax(1).cpu().numpy())
            all_labels.extend(lbls.numpy())
            all_scores.extend(probs[:, pos_idx].cpu().numpy())
    metrics = compute_binary_metrics(all_labels, all_preds, all_scores, class_names, POSITIVE_CLASS)
    return model, metrics


# ==========================================
# RUN THE ABLATION — reports whatever actually happens, A vs B vs C
# ==========================================
print(f"\nRunning ablation with {best_arch} across Conditions A / B / C "
      f"({EPOCHS} epochs each, same split)...\n")
ablation_rows = []
condition_c_model = None
for cond_key, cond_label in [("A_raw", "A_raw"), ("B_static", "B_static"), ("C_adaptive", "C_adaptive")]:
    s_train = samples_for_condition(cond_key, train_idx)
    s_val = samples_for_condition(cond_key, val_idx)
    model, metrics = train_and_eval(s_train, s_val, best_arch)
    metrics["condition"] = cond_label
    ablation_rows.append(metrics)
    printable = {k: round(v, 3) for k, v in metrics.items() if isinstance(v, (int, float))}
    print(f"  {cond_label}: {printable}")
    if cond_key == "C_adaptive":
        condition_c_model = model  # reused below for Grad-CAM: trained on train split only

ablation_df = pd.DataFrame(ablation_rows)
metric_cols = ["accuracy", "precision", "recall", "f1", "roc_auc", "sensitivity", "specificity"]
ranked = ablation_df.sort_values("f1", ascending=False)["condition"].tolist()
print(f"\nAblation ranking by F1 (best to worst): {ranked}")
print("Reported as measured -- if A or B outperforms C here, that's the honest "
      "result to report, not something to re-run until C wins.")

ablation_md_path = os.path.join(REPORT_DIR, "ablation_summary.md")
with open(ablation_md_path, "w") as f:
    f.write("# Ablation Study: Preprocessing Condition Comparison\n\n")
    f.write(f"Architecture: {best_arch}. Same {len(train_idx)}/{len(val_idx)} "
            f"train/val split (grouped, stratified) used for all three conditions.\n\n")
    f.write(ablation_df[["condition"] + metric_cols].round(4).to_markdown(index=False))
    f.write(f"\n\nRanking by F1: {ranked}\n")
print(f"Written: {ablation_md_path}")

# ==========================================
# GRAD-CAM on Condition-C VAL images (unseen during that model's training)
# ==========================================
print("\nRunning Grad-CAM on held-out Condition C validation images...")
target_layer = get_target_layer(condition_c_model, best_arch)
gradcam = GradCAM(condition_c_model, target_layer)
condition_c_model.eval()

val_samples_c = samples_for_condition("C_adaptive", val_idx)
gradcam_rows = []
centroids_in_center_third = 0

for filepath, true_label in tqdm(val_samples_c, desc="Generating Grad-CAMs"):
    img_bgr = cv2.imread(filepath)
    if img_bgr is None:
        continue
    img_rgb = cv2.cvtColor(img_bgr, cv2.COLOR_BGR2RGB)
    input_tensor = tf(Image.fromarray(img_rgb)).unsqueeze(0).to(device)

    cam, pred_idx, confidence = gradcam.generate(input_tensor)
    region, (cx, cy) = describe_focus_region(cam)
    in_center_third = (1 / 3 <= cx <= 2 / 3) and (1 / 3 <= cy <= 2 / 3)
    centroids_in_center_third += int(in_center_third)

    overlay = overlay_heatmap(img_bgr, cam)
    rel_class = class_names[true_label]
    fn = os.path.basename(filepath)
    heatmap_rel = os.path.join("heatmaps", f"heatmap_{rel_class}_{fn}")
    cv2.imwrite(os.path.join(REPORT_DIR, heatmap_rel), overlay)

    gradcam_rows.append({
        "filename": fn,
        "true_class": rel_class,
        "predicted_class": class_names[pred_idx],
        "correct": int(pred_idx == true_label),
        "confidence": round(confidence, 4),
        "focus_region": region,
        "centroid_x_frac": round(cx, 3),
        "centroid_y_frac": round(cy, 3),
        "centroid_in_center_third": in_center_third,
        "path_condition_C": os.path.relpath(filepath, REPORT_DIR),
        "heatmap_path": heatmap_rel,
    })

gradcam_df = pd.DataFrame(gradcam_rows)
pct_centered = 100 * centroids_in_center_third / max(len(gradcam_df), 1)
print(
    f"\n{pct_centered:.1f}% of held-out val images have their Grad-CAM centroid "
    f"in the central third of the frame (where the cardiac chambers typically sit "
    f"after sector masking). This is a measurement, not a guarantee -- inspect "
    f"the actual heatmap_*.png overlays in {HEATMAP_DIR}/ yourself, especially "
    f"any true CHD cases the model got wrong, before citing this figure anywhere."
)

# ==========================================
# MASTER EXCEL REPORT (Arial, navy headers, relative paths, formulas)
# ==========================================
def write_formatted_sheet(ws, df, formula_cols=None):
    for col_idx, header in enumerate(df.columns, 1):
        c = ws.cell(row=1, column=col_idx, value=header)
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
        c.alignment = Alignment(horizontal="center")
    for r, row in enumerate(df.itertuples(index=False), start=2):
        for col_idx, val in enumerate(row, 1):
            ws.cell(row=r, column=col_idx, value=val).font = Font(name="Arial")
    for col_idx in range(1, len(df.columns) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 16

    if formula_cols:
        last_row = len(df) + 1
        avg_row = last_row + 1
        ws.cell(row=avg_row, column=1, value="AVERAGE").font = Font(name="Arial", bold=True)
        for col_name in formula_cols:
            col_idx = list(df.columns).index(col_name) + 1
            col_letter = get_column_letter(col_idx)
            formula = f"=AVERAGE({col_letter}2:{col_letter}{last_row})"
            ws.cell(row=avg_row, column=col_idx, value=formula).font = Font(name="Arial", bold=True)


with pd.ExcelWriter(MASTER_REPORT_PATH, engine="openpyxl") as writer:
    ablation_out = ablation_df[["condition"] + metric_cols].round(4)
    ablation_out.to_excel(writer, sheet_name="ablation_summary", index=False)
    gradcam_df.to_excel(writer, sheet_name="gradcam_results", index=False)
    write_formatted_sheet(writer.sheets["ablation_summary"], ablation_out)
    write_formatted_sheet(writer.sheets["gradcam_results"], gradcam_df, formula_cols=["confidence", "correct"])

print(f"\nMaster report written to {MASTER_REPORT_PATH}")
print(
    "\nBefore using any of this in a paper: the 'centroid in center third' figure "
    "and the ablation ranking are only as trustworthy as the confound diagnostic "
    "from script 1 came back clean. If that diagnostic flagged anything, treat "
    "these numbers as describing what the model learned to shortcut on, not as "
    "evidence it detects CHD."
)
